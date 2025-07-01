import openpyxl
from docx import Document

def extract_numbered_hierarchy_from_docx(docx_path):
    doc = Document(docx_path)
    rows = []
    hierarchy = []

    for para in doc.paragraphs:
        p = para._p  # low-level XML element
        numPr = p.find('.//w:numPr', namespaces=p.nsmap)
        if numPr is not None:
            ilvl = numPr.find('.//w:ilvl', namespaces=p.nsmap)
            if ilvl is not None:
                level = int(ilvl.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val'))
            else:
                level = 0
            text = para.text.strip()
            if len(hierarchy) <= level:
                hierarchy += [''] * (level - len(hierarchy) + 1)
            hierarchy[level] = text
            hierarchy = hierarchy[:level + 1]
            rows.append(hierarchy.copy())
        else:
            # Paragraph without numbering, treat as level 0 or skip
            text = para.text.strip()
            if text:
                hierarchy = [text]
                rows.append(hierarchy.copy())

    return rows

def write_rows_to_excel(rows, output_file):
    if not rows:
        raise ValueError("⚠️ No valid hierarchy items found.")

    max_depth = max(len(row) for row in rows)
    wb = openpyxl.Workbook()
    ws = wb.active

    ws.append([f"H{i+1}" for i in range(max_depth)])

    for row in rows:
        row += [''] * (max_depth - len(row))
        ws.append(row)

    wb.save(output_file)
    print(f"✅ Excel file saved as: {output_file}")

# === CONFIGURATION ===
docx_file = "docx_ex.docx"  # Replace with your actual file name
excel_file = "structured_output.xlsx"

# === RUN ===
rows = extract_numbered_hierarchy_from_docx(docx_file)
write_rows_to_excel(rows, excel_file)
