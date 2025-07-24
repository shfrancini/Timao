import openpyxl
import sys
from docx import Document

def extract_numbered_hierarchy_and_cctp(docx_path):
    doc = Document(docx_path)
    rows = []
    hierarchy = []

    paragraphs = doc.paragraphs
    total_paras = len(paragraphs)

    i = 0
    while i < total_paras:
        para = paragraphs[i]
        p = para._p  # low-level lxml element
        numPr = p.find('.//w:numPr', namespaces=p.nsmap)
        if numPr is not None:
            # This is a numbered paragraph
            ilvl = numPr.find('.//w:ilvl', namespaces=p.nsmap)
            if ilvl is not None:
                level = int(ilvl.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val'))
            else:
                level = 0
            text = para.text.strip()
            if len(hierarchy) <= level:
                hierarchy += [''] * (level - len(hierarchy) + 1)
            hierarchy[level] = text
            hierarchy = hierarchy[:level + 1]

            # Now, look ahead to collect the cctp text block after this numbered paragraph
            cctp_texts = []
            j = i + 1
            while j < total_paras:
                next_p = paragraphs[j]
                next_p_xml = next_p._p
                next_numPr = next_p_xml.find('.//w:numPr', namespaces=next_p_xml.nsmap)
                if next_numPr is not None:
                    # Next numbered paragraph found - stop collecting cctp
                    break
                # Otherwise, add this non-numbered paragraph text if non-empty
                next_text = next_p.text.strip()
                if next_text:
                    cctp_texts.append(next_text)
                j += 1

            cctp_block = "\n".join(cctp_texts)
            rows.append((hierarchy.copy(), cctp_block))

            i = j  # Skip ahead past the collected cctp paragraphs
        else:
            # Not a numbered paragraph — just move on
            i += 1

    return rows

def filter_leaf_paths_with_cctp(rows):
    leaf_rows = []
    for i, (row, cctp) in enumerate(rows):
        row_trimmed = [x for x in row if x]

        is_leaf = True
        for j, (other, _) in enumerate(rows):
            if i == j:
                continue
            other_trimmed = [x for x in other if x]
            if len(other_trimmed) > len(row_trimmed) and other_trimmed[:len(row_trimmed)] == row_trimmed:
                is_leaf = False
                break
        if is_leaf:
            leaf_rows.append((row, cctp))
    return leaf_rows

def write_rows_with_cctp_to_excel(rows, output_file):
    if not rows:
        raise ValueError("⚠️ No valid hierarchy items found.")

    max_depth = max(len(row) for row, _ in rows)
    wb = openpyxl.Workbook()
    ws = wb.active

    # Header: H1, H2, ..., plus CCTP column
    ws.append([f"H{i+1}" for i in range(max_depth)] + ["CCTP"])

    for row, cctp in rows:
        row += [''] * (max_depth - len(row))
        ws.append(row + [cctp])

    wb.save(output_file)
    print(f"✅ Excel file saved as: {output_file}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("❌ Usage: python3 docx_to_xlsx.py <input.docx> <output.xlsx>")
        sys.exit(1)

    docx_file = sys.argv[1]
    excel_file = sys.argv[2]

    rows = extract_numbered_hierarchy_and_cctp(docx_file)
    leaf_rows = filter_leaf_paths_with_cctp(rows)
    write_rows_with_cctp_to_excel(leaf_rows, excel_file)

